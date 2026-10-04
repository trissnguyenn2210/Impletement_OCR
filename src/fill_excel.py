# -*- coding: utf-8 -*-
"""STAGE 4 - doc moi file JSON trong work/results/ va DIEN vao docs/experiment.xlsx.

    python fill_excel.py                 # dien vao ban goc
    python fill_excel.py --dry-run       # chi in ra se dien gi, KHONG ghi file
    python fill_excel.py --out ../docs/experiment_filled.xlsx

An toan:
  - Chi ghi vao DUNG nhung o da map o duoi. Khong cham nhan, tieu de, mau, merge.
  - O nao chua co ket qua thi GIU NGUYEN (khong xoa thanh trong).
  - Gia tri None (khong ap dung) -> ghi chuoi "N/A", khong ghi 0.
    Ghi 0 se bi doc la "lam duoc nhung te", khac han voi "khong lam duoc".
  - Neu file xlsx dang mo trong Excel -> PermissionError, script bao ro va thoat 2.

Ten file ket qua phai dung dang  {method}__{train_on}__{test_on}.json
(run_eval.py tu sinh dung dang nay).
"""
import argparse
import json
import sys
from pathlib import Path

import cfg

# ===================================================================== MAIN
# Bang 2 - cross-domain. 15 dong, thu tu A / C / D / B dung nhu trong file.
# (method_key, train_on, test_on) -> so dong
MAIN_ROWS = {
    ("A", "d2", "d2_val"): 6,
    ("A", "d2", "d1_all"): 7,
    ("A", "d2", "d1_val"): 8,
    ("A", "d1", "d1_val"): 9,
    ("C", "d2", "d2_val"): 10,
    ("C", "d2", "d1_all"): 11,
    ("C", "d2", "d1_val"): 12,
    ("C", "d1", "d1_val"): 13,
    ("D", "d2", "d2_val"): 14,
    ("D", "d2", "d1_all"): 15,
    ("D", "d2", "d1_val"): 16,
    ("D", "d1", "d1_val"): 17,
    ("B", None, "d2_val"): 18,
    ("B", None, "d1_all"): 19,
    ("B", None, "d1_val"): 20,
}

MAIN_COLS = {
    "params_M": 5, "gflops": 6,
    "layout_mAP50": 7, "layout_mAP50_95": 8, "rec_acc": 9, "rec_ned": 10,
    "MFS": 11, "merchant": 12, "address": 13, "date": 14, "total": 15,
    "halluc": 16, "fps": 17,
    "delta_mfs": 18,
    "item_F1": 19, "name_anls": 20, "qty": 21, "price": 22, "amount": 23,
}

# ===================================================================== VLM
# Bang 1 - so sanh cac VLM. Hai block: Train (dong 6-10), Val (dong 15-19).
VLM_MODEL_ROWS = {          # method_key -> (dong block Train, dong block Val)
    "B_paddleocr_vl": (6, 15),
    "B_qwen3vl_2b":   (7, 16),
    "B_qwen3vl_4b":   (8, 17),
    "B_qwen3vl_8b":   (9, 18),
    "B_internvl35_2b": (10, 19),
}
# D1 (MC-OCR) chi co 4 truong header + halluc + fps, KHONG co item
VLM_COLS_D1 = {"MFS": 5, "merchant": 6, "address": 7, "date": 8, "total": 9,
               "halluc": 10, "fps": 11}
# D2 (v3) co them line_items
VLM_COLS_D2 = {"MFS": 12, "merchant": 13, "address": 14, "date": 15, "total": 16,
               "item_F1": 17, "name_anls": 18, "qty": 19, "price": 20,
               "amount": 21, "halluc": 22, "fps": 23}
# test_on -> (so dong trong cap, map cot)
VLM_CELL = {
    "d1_train": (0, VLM_COLS_D1),
    "d1_val":   (1, VLM_COLS_D1),
    "d2_train": (0, VLM_COLS_D2),
    "d2_val":   (1, VLM_COLS_D2),
}

# ================================================================== quality
# Bang 3 - phan tang theo chat luong anh. CHI tren D1 (chi D1 co anno_image_quality).
QUALITY_ROWS = {"A": 6, "B": 7, "C": 8, "D": 9}
QUALITY_BIN_COL = {"clear": 3, "mid": 8, "blur": 13}   # moi bin: MFS, 4 truong
QUALITY_FIELD_ORDER = ("merchant", "address", "date", "total")
QUALITY_DELTA_COL = 18
# dong nao cua bang quality lay tu ket qua nao
QUALITY_SOURCE = {"A": ("A", "d1", "d1_val"), "B": ("B", None, "d1_val"),
                  "C": ("C", "d1", "d1_val"), "D": ("D", "d1", "d1_val")}

NA = "N/A"


# ------------------------------------------------------------------ helpers
def r1(v, nd=2):
    """None -> 'N/A'. So -> lam tron. Con lai -> giu nguyen."""
    if v is None:
        return NA
    if isinstance(v, (int, float)):
        return round(float(v), nd)
    return v


def load_results() -> dict:
    """Doc moi JSON trong work/results/. Key = (method, train_on, test_on)."""
    out = {}
    if not cfg.RESULT_DIR.exists():
        return out
    for p in sorted(cfg.RESULT_DIR.glob("*.json")):
        try:
            d = json.loads(p.read_text(encoding="utf-8"))
        except Exception as e:
            print(f"   BO QUA {p.name}: khong doc duoc JSON ({e})")
            continue
        key = (d.get("method"), d.get("train_on"), d.get("test_on"))
        if None in (key[0], key[2]):
            print(f"   BO QUA {p.name}: thieu method hoac test_on")
            continue
        out[key] = d
        out[key]["_file"] = p.name
    return out


def flat(res: dict) -> dict:
    """Gop moi con so cua mot lan chay vao mot dict phang de tra cot."""
    agg = res.get("agg") or {}
    items = res.get("items") or {}
    f = dict(agg.get("per_field") or {})
    f["MFS"] = agg.get("MFS")
    f["halluc"] = agg.get("halluc")
    f["fps"] = res.get("fps")
    for k in ("params_M", "gflops", "layout_mAP50", "layout_mAP50_95",
              "rec_acc", "rec_ned"):
        f[k] = res.get(k)
    for k in ("item_F1", "name_anls", "qty", "price", "amount"):
        f[k] = items.get(k) if items else None
    return f


def method_key(m: str) -> str:
    """'A' / 'A_ppstructure' / 'dummy_perfect' -> chu cai phuong phap."""
    return m.split("_")[0]


# -------------------------------------------------------------------- fills
def _pick(results, mk, tr, te):
    """Tim ket qua cho mot dong MAIN.

    PHAI uu tien khop TEN CHINH XAC truoc. Ly do: method_key('B_qwen3vl_4b')
    cung ra 'B', nen neu chi so method_key thi dong B cua bang MAIN co the
    nhan so cua MOT MODEL VLM BAT KY, tuy thu tu file trong thu muc.
    Loi nay khong lam crash - bang van day so, chi la so sai model.
    """
    exact_hit = results.get((mk, tr, te))
    if exact_hit is not None:
        return exact_hit
    cands = [d for (m, t, s), d in results.items()
             if method_key(m) == mk and t == tr and s == te]
    if not cands:
        return None
    if len(cands) > 1:
        names = sorted(d["_file"] for d in cands)
        print(f"   CANH BAO dong {mk}/{tr}/{te}: {len(cands)} file cung khop "
              f"{names}. Lay {names[0]}. Hay chay 'run_eval.py --method {mk}' "
              f"de co ket qua ro rang.")
        cands.sort(key=lambda d: d["_file"])
    return cands[0]


def fill_main(ws, results, log):
    # 1) dien 15 dong
    for (mk, tr, te), row in MAIN_ROWS.items():
        res = _pick(results, mk, tr, te)
        if res is None:
            continue
        f = flat(res)
        for name, col in MAIN_COLS.items():
            if name == "delta_mfs":
                continue
            if name in f:
                ws.cell(row, col).value = r1(f[name])
        log.append(f"   MAIN  dong {row:2d}  {mk} train={tr} test={te}"
                   f"   MFS={r1(f.get('MFS'))}   <- {res['_file']}")

    # 2) delta_mfs = MFS cross-domain - MFS in-domain.
    #    in-domain cua moi phuong phap = dong train d2 / test d2_val.
    for mk in ("A", "C", "D", "B"):
        base_row = MAIN_ROWS.get((mk, "d2" if mk != "B" else None, "d2_val"))
        base = ws.cell(base_row, MAIN_COLS["MFS"]).value if base_row else None
        if not isinstance(base, (int, float)):
            continue
        for (m2, t2, s2), row in MAIN_ROWS.items():
            if m2 != mk or s2 == "d2_val":
                continue
            cur = ws.cell(row, MAIN_COLS["MFS"]).value
            if isinstance(cur, (int, float)):
                ws.cell(row, MAIN_COLS["delta_mfs"]).value = round(cur - base, 2)


def fill_vlm(ws, results, log):
    for (m, tr, te), res in results.items():
        if m not in VLM_MODEL_ROWS or te not in VLM_CELL:
            continue
        idx, colmap = VLM_CELL[te]
        row = VLM_MODEL_ROWS[m][idx]
        f = flat(res)
        for name, col in colmap.items():
            if name in f:
                ws.cell(row, col).value = r1(f[name])
        log.append(f"   VLM   dong {row:2d}  {m} test={te}"
                   f"   MFS={r1(f.get('MFS'))}   <- {res['_file']}")


def fill_quality(ws, results, log):
    for mk, src in QUALITY_SOURCE.items():
        res = _pick(results, *src)
        if res is None or not res.get("by_quality"):
            continue
        row = QUALITY_ROWS[mk]
        bq = res["by_quality"]
        mfs = {}
        for bin_name, c0 in QUALITY_BIN_COL.items():
            sub = bq.get(bin_name)
            if sub is None:
                ws.cell(row, c0).value = NA
                continue
            ws.cell(row, c0).value = r1(sub.get("MFS"))
            mfs[bin_name] = sub.get("MFS")
            pf = sub.get("per_field") or {}
            for k, fld in enumerate(QUALITY_FIELD_ORDER, 1):
                ws.cell(row, c0 + k).value = r1(pf.get(fld))
        if isinstance(mfs.get("clear"), float) and isinstance(mfs.get("blur"), float):
            ws.cell(row, QUALITY_DELTA_COL).value = round(mfs["clear"] - mfs["blur"], 2)
        log.append(f"   qual  dong {row:2d}  {mk}   "
                   f"ro={r1(mfs.get('clear'))} mo={r1(mfs.get('blur'))}"
                   f"   <- {res['_file']}")


# --------------------------------------------------------------------- main
def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--xlsx", default=str(cfg.ROOT / "docs" / "experiment.xlsx"))
    ap.add_argument("--out", default=None, help="ghi ra file khac (giu ban goc)")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    try:
        import openpyxl
    except ImportError:
        print("Thieu openpyxl. Chay:  pip install openpyxl")
        return 2

    src = Path(a.xlsx)
    if not src.exists():
        print(f"Khong thay {src}")
        return 2

    results = load_results()
    print(f"Doc {len(results)} file ket qua tu {cfg.RESULT_DIR}")
    if not results:
        print("Chua co ket qua nao. Chay run_eval.py truoc.")
        return 0
    for k in sorted(results, key=lambda x: (x[0], str(x[1]), x[2])):
        print(f"   co: {k[0]:16s} train={str(k[1]):5s} test={k[2]}")

    wb = openpyxl.load_workbook(src)
    log = []
    print()
    print("DIEN:")
    fill_main(wb["MAIN"], results, log)
    fill_vlm(wb["VLM"], results, log)
    fill_quality(wb["quality"], results, log)
    for line in log:
        print(line)
    if not log:
        print("   (khong co o nao khop. Kiem tra --method / --train-on / --test-on"
              " co dung ten trong MAIN_ROWS / VLM_MODEL_ROWS khong)")

    if a.dry_run:
        print("\n--dry-run: KHONG ghi file.")
        return 0

    dst = Path(a.out) if a.out else src
    try:
        wb.save(dst)
    except PermissionError:
        print(f"\nKHONG GHI DUOC {dst}")
        print("File dang mo trong Excel. Dong Excel roi chay lai,")
        print("hoac dung:  python fill_excel.py --out ../docs/experiment_filled.xlsx")
        return 2
    print(f"\nDa ghi {len(log)} dong vao {dst}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
